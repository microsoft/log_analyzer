# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.
"""Digest sensor-analysis workbooks into one reviewable HTML page.

Sensor-analysis workbooks are large (tens of MB, dozens of sheets) so nobody
opens them. This turns them into a single page that surfaces what actually
matters: status anomalies and sensors whose value changed significantly, with
error-type counters (CRC / fatal / uncorrectable / link-down ...) ranked first.

Inputs, all next to each other in the scan folder (names are configurable):

* ``<base><WORKBOOK_SUFFIX>`` - workbook matched by ``SENSOR_XLSX_GLOB``
  (default ``*sensor*.xlsx``). Its optional ``Summary`` and ``Status_bad``
  sheets are shown as-is.
* ``<workbook stem>.log`` - optional companion log with one line per changed
  sensor, e.g. ``[change] 6 :: PWR_SUM :: first=600 last=2040 delta=1440
  min=120 max=2100`` (tag set by ``CHANGE_LINE_TAG``).
* ``<base>.csv`` - optional raw collector CSV (IPMI SDR style columns
  ``Name,Reading,System-Timestamp[,Status]`` or Redfish telemetry style
  ``MetricProperty,MetricValue,Timestamp[,MetricValueStale]``) used for
  data-quality checks and charts.

Usage:
    python sensor_xlsx_digest.py --scan_folder <log folder> [--output <html>]
        [--xlsx_glob "*sensor*.xlsx"] [--workbook_suffix "_analysis.xlsx"]
"""

import argparse
import html
import logging
import os
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    _common_dir = str(Path(__file__).resolve().parent.parent / "common")
    if _common_dir not in sys.path:
        sys.path.insert(0, _common_dir)
    import brand
except ImportError:  # pragma: no cover - branding is optional
    brand = None

logger = logging.getLogger(__name__)

OUTPUT_NAME = "sensor_xlsx_digest.html"

# Workbooks to digest, relative to the scan folder.
# Override with --xlsx_glob or LOG_ANALYZER_SENSOR_XLSX_GLOB.
SENSOR_XLSX_GLOB = os.getenv("LOG_ANALYZER_SENSOR_XLSX_GLOB", "*sensor*.xlsx")

# Workbook-name suffix that is replaced by ".csv" to find the raw collector CSV
# ("BMC_sensor_reading_analysis.xlsx" -> "BMC_sensor_reading.csv"). Plain
# ".xlsx" is always tried as a fallback. Override with --workbook_suffix or
# LOG_ANALYZER_SENSOR_WORKBOOK_SUFFIX.
WORKBOOK_SUFFIX = os.getenv("LOG_ANALYZER_SENSOR_WORKBOOK_SUFFIX", "_analysis.xlsx")

# Tag that starts a changed-sensor line in the companion .log.
# Override with LOG_ANALYZER_SENSOR_CHANGE_TAG.
CHANGE_LINE_TAG = os.getenv("LOG_ANALYZER_SENSOR_CHANGE_TAG", "change")

# Values that are not numbers but still carry a verdict. A workbook generator
# that coerces every reading to a number turns these states into NaN, so they
# are invisible there - we surface them here instead.
GOOD_STATES = {
    "ok", "enabled", "linkup", "nolinkdown", "up", "normal", "true", "false",
    "none", "standbyspare", "starting", "pt0s", "inactive", "active",
}
PROBLEM_STATES = {
    "training": "link never finished training",
    "linkdown": "link is down",
    "down": "link/device is down",
    "absent": "device is absent",
    "failed": "device reported Failed",
    "critical": "health is Critical",
    "warning": "health is Warning",
    "degraded": "device reported Degraded",
    "error": "device reported Error",
    "unknown": "state is Unknown",
    "unavailableoffline": "device is offline",
    "standbyoffline": "device is offline",
    "sensor failure": "BMC reports a sensor failure",
    "sensorfailure": "BMC reports a sensor failure",
    "fault": "device reported a fault",
    "faulty": "device reported faulty",
    "no communication": "no communication with the device",
}

# Not faults: administratively disabled / unpopulated ports are normal on a
# partially populated board. Surfaced as NOTICE so they stay visible.
NOTICE_STATES = {
    "disabled": "port/device administratively disabled or unpopulated",
    "unavailable": "device not present in this configuration",
}

# The sensor exists but carries no value - a data-quality gap, not a fault.
NOVALUE_STATES = {
    "na": "reading unavailable",
    "n/a": "reading unavailable",
    "nan": "reading unavailable",
    "null": "reading unavailable",
}
# 0xFF / 0xFFFF style sentinels mean "no reading" on IPMI-style sensors.
FF_SENTINEL_RE = re.compile(r"^(?:0x)?(?:ff)+$", re.IGNORECASE)

# Descriptive values that record what happened, not that something is wrong
# (for example how a device was last reset).
BENIGN_INFO_RE = re.compile(
    r"^(?:Conventional|Fundamental|\w*Reset|PowerCycle|Warm|Cold|None|Unspecified)$",
    re.IGNORECASE)

# BMCs serve readings as "45 degrees C" / "0.22 Amps" / "53.32 Volts". Such rows
# typically never parse as a bare number, so treating them as opaque states
# floods the report with healthy temperature sensors.
UNIT_SUFFIX_RE = re.compile(
    r"^\s*([-+]?\d+(?:\.\d+)?)\s*"
    r"(?:degrees?\s*[CF]|deg\s*[CF]|[CF]\b|Volts?|V\b|Amps?|A\b|Watts?|W\b|"
    r"RPM|percent|%|Joules?|J\b|Hz|MHz|GHz|bytes?|KB|MB|GB|seconds?|s\b|ms)\s*$",
    re.IGNORECASE)


def parse_reading(value):
    """Number behind a reading, unit stripped. None if it is a real state string."""
    v = _to_float(value)
    if v is not None:
        return v
    m = UNIT_SUFFIX_RE.match(str(value))
    return float(m.group(1)) if m else None

# A sensor that stops reporting for more than this fraction of the run is
# treated as having gone missing.
MISSING_FRACTION = 0.2

# A state shared by ~all sensors of the same metric leaf is this platform's
# firmware baseline, not a fleet-wide failure. Demote it to NOTICE.
PEER_CONSENSUS = 0.9
PEER_MIN_SENSORS = 5

# Counter-style sensors where any increase is a genuine hardware concern.
ERROR_METRIC_RE = re.compile(
    r"fatal|uncorrect|crcerror|packetcrc|linkdowned|linkerrorrecovery|"
    r"malformed|rxerror|symbolerror|effectiveerror|runtimeerror|trainingerror|"
    r"correctableerror|l0torecovery|bitErrorRate|effectiveber|\bber\b",
    re.IGNORECASE,
)

def change_line_re(tag=None):
    """Regex for "[<tag>] 6 :: PWR_SUM :: first=600 last=2040 delta=1440 min=120 max=2100"."""
    return re.compile(
        r"^\[" + re.escape(tag or CHANGE_LINE_TAG) + r"\]\s+(?P<count>\d+)\s*::\s*(?P<name>.+?)\s*::\s*"
        r"first=(?P<first>\S+)\s+last=(?P<last>\S+)\s+delta=(?P<delta>\S+)\s+"
        r"min=(?P<min>\S+)\s+max=(?P<max>\S+)\s*$"
    )


def _to_float(text):
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def raw_csv_for(xlsx_path, workbook_suffix=None):
    """`BMC_sensor_reading_analysis.xlsx` -> `BMC_sensor_reading.csv` (if present)."""
    name = xlsx_path.name
    for suffix in (workbook_suffix or WORKBOOK_SUFFIX, ".xlsx"):
        if suffix and name.lower().endswith(suffix.lower()):
            candidate = xlsx_path.with_name(name[:-len(suffix)] + ".csv")
            if candidate.is_file():
                return candidate
    return None


def load_raw_csv(csv_path):
    """Load the collector CSV keeping the columns a workbook generator discards.

    Workbook generators typically keep only name/value/timestamp and coerce the
    value to a number, so stale flags and textual states such as "Training"
    never reach the workbook. We keep them.
    """
    try:
        import pandas as pd
    except ImportError:
        logger.warning("pandas not available - skipping data-quality checks")
        return None

    try:
        df = pd.read_csv(csv_path, dtype=str, keep_default_na=False,
                         low_memory=False, on_bad_lines="skip")
    except Exception as exc:
        logger.warning("Could not read %s: %s", csv_path.name, exc)
        return None

    cols = set(df.columns)
    if {"Name", "Reading", "System-Timestamp"} <= cols:          # IPMI SDR layout
        out = pd.DataFrame({
            "MetricProperty": df["Name"],
            "MetricValue": df["Reading"],
            "Timestamp": df["System-Timestamp"],
            "RawStatus": df.get("Status", ""),
            "Stale": "",
        })
    elif {"MetricProperty", "MetricValue", "Timestamp"} <= cols:  # Redfish telemetry layout
        out = pd.DataFrame({
            "MetricProperty": df["MetricProperty"],
            "MetricValue": df["MetricValue"],
            "Timestamp": df["Timestamp"],
            "RawStatus": df.get("RawStatus", ""),
            "Stale": df.get("MetricValueStale", ""),
        })
    else:
        logger.warning("Unrecognised CSV layout: %s", csv_path.name)
        return None

    out["Timestamp"] = pd.to_datetime(out["Timestamp"], errors="coerce")
    return out.dropna(subset=["Timestamp"])


def _state_verdict(value):
    """Classify a reading that is not a bare number. None = nothing to report."""
    text = str(value).strip()
    if text == "":
        return "EMPTY", "no reading returned"
    if parse_reading(text) is not None:
        return None                       # "45 degrees C" is a healthy reading
    lowered = text.lower()
    # ACPI power state, e.g. "S0/G0: working" - informational, and S0 means the
    # host is actually running.
    if lowered.startswith(("s0/g0", "s0 ", "working")):
        return None
    # "Conventional" / "...Reset" etc. simply record how a device last reset.
    if BENIGN_INFO_RE.search(text):
        return None
    if FF_SENTINEL_RE.match(text):
        return "WARNING", "0xFF sentinel - sensor reports no reading"
    if lowered in GOOD_STATES:
        return None
    reason = PROBLEM_STATES.get(lowered)
    if reason:
        return "PROBLEM", reason
    reason = NOVALUE_STATES.get(lowered)
    if reason:
        return "WARNING", reason
    reason = NOTICE_STATES.get(lowered)
    if reason:
        return "NOTICE", reason
    return "INFO", "non-numeric state (not recognised)"


_DIGITS_RE = re.compile(r"\d+")


def _leaf_of(sensor_name):
    """Peer group key: the full path with indices normalised away.

    '.../Ports/ProcessorModule_0_PCIeBus_0_PCIeLink_1#/Status/State' and its
    siblings collapse to one key, while GPU/memory State sensors stay separate.
    Grouping on the bare leaf ('Status/State') would lump unrelated device
    classes together and hide real outliers.
    """
    return _DIGITS_RE.sub("N", str(sensor_name))


def _leaf_label(sensor_name):
    text = str(sensor_name)
    return text.rsplit("#/", 1)[-1] if "#/" in text else text.rsplit("/", 1)[-1]


def detect_active_window(df, min_gap_minutes=15, gap_factor=10):
    """Restrict analysis to the last continuous run, cutting power-cycle noise.

    Everything is derived from the CSV's own UTC timestamps, so there is no
    local-time vs UTC skew to get wrong (host logs often print local time while
    the CSVs print UTC).

    A sensor CSV can span a whole day yet carry long dead periods left over
    from earlier setup phases, while the actual test ran only a few hours.
    Sampling simply stops while the host is down, so a gap far larger than the
    normal cadence marks a power cycle.
    """
    import pandas as pd

    if df is None or df.empty:
        return None
    instants = pd.Series(sorted(df["Timestamp"].unique()))
    if len(instants) < 3:
        return None
    gaps = instants.diff().dt.total_seconds().dropna()
    cadence = gaps.median()
    threshold = max(min_gap_minutes * 60, cadence * gap_factor)
    breaks = [i for i in gaps.index if gaps[i] > threshold]
    start_idx = breaks[-1] if breaks else 0
    start, end = instants.iloc[start_idx], instants.iloc[-1]
    return {
        "start": start,
        "end": end,
        "cadence_s": cadence,
        "threshold_s": threshold,
        "sessions": len(breaks) + 1,
        "dropped": int((df["Timestamp"] < start).sum()),
        "total": len(df),
    }


def check_data_quality(df):
    """Find stale readings, bad states, null/FF readings and missing sensors.

    These never show up in a numeric workbook: every non-numeric reading is
    coerced to NaN there, and MetricValueStale is typically not read.
    """
    import pandas as pd

    findings = []
    if df is None or df.empty:
        return findings

    numeric = df["MetricValue"].map(parse_reading)
    non_numeric = df[numeric.isna()]

    # 1) Textual states - report per sensor, judged on the LAST value seen.
    if not non_numeric.empty:
        ordered = non_numeric.sort_values("Timestamp")
        groups = list(ordered.groupby("MetricProperty", sort=False))

        # Peer consensus: if nearly every sensor of the same metric leaf ends in
        # the same state, that is the platform/firmware baseline for this box,
        # not 140 simultaneous failures. Demote it so the real outliers show.
        peers, ends = {}, {}
        for name, grp in groups:
            leaf = _leaf_of(name)
            last = str(grp["MetricValue"].iloc[-1])
            ends[name] = (leaf, last)
            peers.setdefault(leaf, []).append(last)
        baseline = {}
        for leaf, values in peers.items():
            if len(values) < PEER_MIN_SENSORS:
                continue
            top, hits = Counter(values).most_common(1)[0]
            if hits / len(values) >= PEER_CONSENSUS:
                baseline[leaf] = (top, hits, len(values))

        empty_sensors = []
        for name, grp in groups:
            verdicts = {}
            for value in grp["MetricValue"].unique():
                v = _state_verdict(value)
                if v:
                    verdicts[str(value)] = v
            if not verdicts:
                continue
            leaf, last_value = ends[name]
            last_verdict = _state_verdict(last_value)

            # A blank reading is a poll that returned nothing - a collection gap,
            # not a sensor state. Roll them into one line instead of one per
            # sensor, otherwise every intermittent temperature probe gets its
            # own row.
            if set(v[0] for v in verdicts.values()) == {"EMPTY"}:
                empty_sensors.append((name, len(grp), last_value))
                continue

            if last_verdict and last_verdict[0] == "PROBLEM":
                agreed = baseline.get(leaf)
                if agreed and agreed[0] == last_value:
                    sev = "NOTICE"
                    why = (f"'{last_value}' on {agreed[1]}/{agreed[2]} identical peer "
                           f"{_leaf_label(name)} sensors - baseline, not an outlier")
                else:
                    sev = "CRITICAL"
                    why = f"ends the run in '{last_value}' - {last_verdict[1]}"
                    if agreed:
                        why += f" (its {agreed[2]} peers report '{agreed[0]}')"
            elif last_verdict and last_verdict[0] == "EMPTY":
                # The last poll came back blank, but do not let that hide a real
                # state such as "Sensor failure" seen earlier in the run.
                worst = next((v for v in verdicts.values() if v[0] == "PROBLEM"), None)
                if worst:
                    hit = next(k for k, v in verdicts.items() if v[0] == "PROBLEM")
                    sev = "CRITICAL"
                    why = (f"reported '{hit}' during the run - {worst[1]}; "
                           f"the final poll returned no reading")
                else:
                    sev = "WARNING"
                    why = "last poll returned no reading"
            elif last_verdict and last_verdict[0] in ("WARNING", "NOTICE"):
                sev = last_verdict[0]
                why = f"ends the run in '{last_value}' - {last_verdict[1]}"
            elif any(v[0] == "PROBLEM" for v in verdicts.values()):
                hit = next(k for k, v in verdicts.items() if v[0] == "PROBLEM")
                sev, why = "WARNING", f"was '{hit}' during the run, ended '{last_value}'"
            elif any(v[0] == "WARNING" for v in verdicts.values()):
                sev = "WARNING"
                why = next(v[1] for v in verdicts.values() if v[0] == "WARNING")
            elif any(v[0] == "NOTICE" for v in verdicts.values()):
                sev = "NOTICE"
                why = next(v[1] for v in verdicts.values() if v[0] == "NOTICE")
            else:
                sev, why = "INFO", next(iter(verdicts.values()))[1]

            shown = [v if v.strip() else "(empty)" for v in sorted(verdicts)]
            findings.append({
                "kind": "State / non-numeric reading",
                "sensor": name,
                "detail": ", ".join(shown) or last_value,
                "last": last_value if last_value.strip() else "(empty)",
                "count": len(grp),
                "severity": sev,
                "why": why,
            })

        if empty_sensors:
            total_blank = sum(c for _, c, _ in empty_sensors)
            sample = ", ".join(n for n, _, _ in empty_sensors[:4])
            findings.append({
                "kind": "Blank readings",
                "sensor": f"{len(empty_sensors)} sensors ({sample}"
                          f"{', ...' if len(empty_sensors) > 4 else ''})",
                "detail": "(empty)",
                "last": "-",
                "count": total_blank,
                "severity": "NOTICE",
                "why": "some polls returned an empty value; every other sample "
                       "for these sensors was a valid reading",
            })

    # 2) Stale flag - present in the CSV but dropped before the workbook.
    stale = df[df["Stale"].astype(str).str.lower() == "true"]
    if not stale.empty:
        for name, grp in stale.groupby("MetricProperty", sort=False):
            findings.append({
                "kind": "Stale reading",
                "sensor": name,
                "detail": "MetricValueStale=True",
                "last": str(grp["MetricValue"].iloc[-1]),
                "count": len(grp),
                "severity": "WARNING",
                "why": "controller served a cached value - sensor did not refresh",
            })

    # 3) Sensors that stop reporting part-way through the run.
    t_min, t_max = df["Timestamp"].min(), df["Timestamp"].max()
    span = (t_max - t_min).total_seconds()
    if span > 0:
        last_seen = df.groupby("MetricProperty")["Timestamp"].max()
        cutoff = span * MISSING_FRACTION
        for name, seen in last_seen.items():
            gap = (t_max - seen).total_seconds()
            if gap > cutoff:
                findings.append({
                    "kind": "Sensor stopped reporting",
                    "sensor": name,
                    "detail": f"last seen {seen}",
                    "last": "-",
                    "count": 0,
                    "severity": "CRITICAL",
                    "why": f"absent for the final {gap / 3600:.1f} h of a {span / 3600:.1f} h run",
                })

    findings.sort(key=lambda f: (SEVERITY_ORDER.get(f["severity"], 9), -f["count"], f["sensor"]))
    return findings


def series_for(df, sensor_name, max_points=400):
    """Numeric time series for one sensor, thinned for plotting."""
    import pandas as pd

    if df is None or df.empty:
        return []
    sub = df[df["MetricProperty"] == sensor_name]
    if sub.empty:
        return []
    sub = sub.sort_values("Timestamp")
    values = sub["MetricValue"].map(parse_reading)
    pairs = [(t, v) for t, v in zip(sub["Timestamp"], values) if v is not None]
    if len(pairs) > max_points:
        step = len(pairs) / max_points
        pairs = [pairs[int(i * step)] for i in range(max_points)]
    return pairs


def state_series_for(df, sensor_name, max_points=600):
    """Raw (timestamp, state-string) series - for sensors that never report numbers."""
    if df is None or df.empty:
        return []
    sub = df[df["MetricProperty"] == sensor_name]
    if sub.empty:
        return []
    sub = sub.sort_values("Timestamp")
    pairs = list(zip(sub["Timestamp"], sub["MetricValue"].astype(str)))
    if len(pairs) > max_points:
        step = len(pairs) / max_points
        pairs = [pairs[int(i * step)] for i in range(max_points)]
    return pairs


def _state_color(value):
    verdict = _state_verdict(value)
    if verdict is None:
        return "#2e7d32"                      # healthy
    return {"PROBLEM": "#c62828", "WARNING": "#ef6c00",
            "NOTICE": "#1565c0"}.get(verdict[0], "#9e9e9e")


def svg_state_timeline(pairs, width=760, height=76):
    """Horizontal colour band: what state the sensor held, start to end.

    A line chart cannot show LinkStatus - it is never numeric - yet those are the
    findings a user most wants to trace from beginning to end.
    """
    if not pairs:
        return '<div class="nochart">no samples for this sensor</div>'

    pad_l, pad_r, pad_t = 8, 8, 8
    band_h, plot_w = 26, width - pad_l - pad_r
    n = len(pairs)
    segments, legend, seen = [], [], {}
    run_start, run_value = 0, pairs[0][1]

    def _flush(start, end, value):
        x = pad_l + plot_w * start / n
        w = max(1.0, plot_w * (end - start) / n)
        color = _state_color(value)
        seen.setdefault(value, color)
        segments.append(
            f'<rect x="{x:.1f}" y="{pad_t}" width="{w:.1f}" height="{band_h}" '
            f'fill="{color}"><title>{_esc(value)} ({end - start} samples)</title></rect>'
        )

    for i in range(1, n):
        if pairs[i][1] != run_value:
            _flush(run_start, i, run_value)
            run_start, run_value = i, pairs[i][1]
    _flush(run_start, n, run_value)

    lx = pad_l
    for value, color in seen.items():
        legend.append(
            f'<rect x="{lx}" y="{pad_t + band_h + 10}" width="10" height="10" fill="{color}"/>'
            f'<text x="{lx + 14}" y="{pad_t + band_h + 19}" font-size="10" fill="#555">'
            f'{_esc(value)}</text>'
        )
        lx += 22 + 7 * len(str(value))

    t0 = pairs[0][0].strftime("%m-%d %H:%M")
    t1 = pairs[-1][0].strftime("%m-%d %H:%M")
    return f"""<svg class="spark" viewBox="0 0 {width} {height}" width="100%" height="{height}"
 preserveAspectRatio="none" role="img">
{''.join(segments)}
{''.join(legend)}
<text x="{pad_l}" y="{height - 4}" font-size="10" fill="#777">{t0}</text>
<text x="{width - pad_r}" y="{height - 4}" text-anchor="end" font-size="10" fill="#777">{t1}</text>
</svg>"""


def svg_chart(pairs, width=760, height=150):
    """Inline SVG line chart - no JS, no external files, survives copy/paste."""
    if len(pairs) < 2:
        return '<div class="nochart">not enough numeric samples to plot</div>'

    values = [v for _, v in pairs]
    lo, hi = min(values), max(values)
    span = (hi - lo) or (abs(hi) or 1.0)
    pad_l, pad_b, pad_t, pad_r = 78, 26, 12, 12
    plot_w = width - pad_l - pad_r
    plot_h = height - pad_t - pad_b

    def _x(i):
        return pad_l + (plot_w * i / (len(pairs) - 1))

    def _y(v):
        return pad_t + plot_h - ((v - lo) / span) * plot_h

    points = " ".join(f"{_x(i):.1f},{_y(v):.1f}" for i, v in enumerate(values))
    t0 = pairs[0][0].strftime("%m-%d %H:%M")
    t1 = pairs[-1][0].strftime("%m-%d %H:%M")

    def _fmt(v):
        return f"{v:.6g}"

    return f"""<svg class="spark" viewBox="0 0 {width} {height}" width="100%" height="{height}"
 preserveAspectRatio="none" role="img">
<rect x="{pad_l}" y="{pad_t}" width="{plot_w}" height="{plot_h}" fill="#fbfbfd" stroke="#e0e0e0"/>
<polyline points="{points}" fill="none" stroke="#c62828" stroke-width="1.6"/>
<circle cx="{_x(0):.1f}" cy="{_y(values[0]):.1f}" r="3" fill="#2e7d32"/>
<circle cx="{_x(len(values) - 1):.1f}" cy="{_y(values[-1]):.1f}" r="3" fill="#c62828"/>
<text x="{pad_l - 6}" y="{pad_t + 9}" text-anchor="end" font-size="10" fill="#555">{_fmt(hi)}</text>
<text x="{pad_l - 6}" y="{pad_t + plot_h}" text-anchor="end" font-size="10" fill="#555">{_fmt(lo)}</text>
<text x="{pad_l}" y="{height - 8}" font-size="10" fill="#777">{t0} (start {_fmt(values[0])})</text>
<text x="{width - pad_r}" y="{height - 8}" text-anchor="end" font-size="10" fill="#777">
{_fmt(values[-1])} end) {t1}</text>
</svg>"""


def classify(name, delta, maximum):
    """Rank a changed-sensor entry. Counters that grew are the real findings."""
    is_error_metric = bool(ERROR_METRIC_RE.search(name))
    delta_val = _to_float(delta)
    max_val = _to_float(maximum)

    if is_error_metric:
        if delta_val is not None and delta_val > 0:
            return "CRITICAL", "error counter increased during the run"
        if max_val is not None and max_val > 0:
            return "WARNING", "non-zero error counter (no growth in this run)"
        return "INFO", "error counter stayed at zero"

    if delta_val is not None and delta_val != 0:
        return "NOTICE", "value changed between start and end"
    return "INFO", "fluctuated but returned to its starting value"


SEVERITY_ORDER = {"CRITICAL": 0, "WARNING": 1, "NOTICE": 2, "INFO": 3}


def looks_like_sentinel(value):
    """True for large integers whose low 16 bits are all zero.

    Some firmware reports error counters such as 0x201a0000 / 0x20450000 on
    every port. A real error counter does not land on a 64K boundary on dozens
    of ports at once - these are unimplemented or mis-encoded firmware fields.
    """
    v = _to_float(value)
    if v is None or v == 0 or v != int(v):
        return False
    n = abs(int(v))
    return n >= (1 << 24) and (n & 0xFFFF) == 0


def demote_changes_by_peers(entries, raw_df):
    """Demote numeric findings that every identical peer sensor also reports.

    check_data_quality() already does this for textual states. Without the same
    treatment here, dozens of ports reporting one identical firmware constant
    all show up as CRITICAL and bury the genuine outliers.
    """
    if raw_df is None or raw_df.empty or not entries:
        return entries

    last_by_sensor = (raw_df.sort_values("Timestamp")
                            .groupby("MetricProperty")["MetricValue"].last())
    groups = {}
    for sensor, last in last_by_sensor.items():
        groups.setdefault(_leaf_of(sensor), []).append(str(last))

    for e in entries:
        if e["severity"] != "CRITICAL":
            continue
        # The change log stores a truncated "5.38575e+08"; the CSV holds the exact
        # 538574848. Always judge on the CSV value or both the peer match and the
        # sentinel test silently fail.
        exact = last_by_sensor.get(e["name"])
        exact = str(exact) if exact is not None else str(e["last"])
        peers = groups.get(_leaf_of(e["name"]))
        sentinel = looks_like_sentinel(exact)
        if not peers or len(peers) < PEER_MIN_SENSORS:
            if sentinel:
                e["severity"] = "WARNING"
                e["why"] = (f"value {exact} has all-zero low 16 bits - "
                            f"looks like an unimplemented firmware field, not a real count")
            continue
        top, hits = Counter(peers).most_common(1)[0]
        share = hits / len(peers)
        top_val, exact_val = _to_float(top), _to_float(exact)
        same = (top == exact) or (
            top_val is not None and exact_val is not None
            and abs(top_val - exact_val) <= abs(exact_val) * 1e-6
        )
        if same and share >= PEER_CONSENSUS:
            e["severity"] = "NOTICE"
            e["why"] = (f"{hits}/{len(peers)} identical peer sensors end on the very "
                        f"same value {exact} - platform constant, not a fault")
            if sentinel:
                e["why"] += " (low 16 bits all zero: unimplemented firmware field)"
        elif sentinel:
            e["severity"] = "WARNING"
            e["why"] = (f"differs from its {len(peers)} peers, but {exact} has all-zero "
                        f"low 16 bits - unreliable firmware field, verify before acting")

    entries.sort(key=lambda e: (SEVERITY_ORDER[e["severity"]], -int(e["count"]), e["name"]))
    return entries


SEVERITY_COLOR = {
    "CRITICAL": "#c62828",
    "WARNING": "#ef6c00",
    "NOTICE": "#1565c0",
    "INFO": "#666666",
}


def read_summary(xlsx_path):
    """Read the Summary and Status_bad sheets without loading the whole book."""
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl not available - skipping %s", xlsx_path.name)
        return None

    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Could not open %s: %s", xlsx_path.name, exc)
        return None

    info = {"file": xlsx_path.name, "summary": [], "status_bad": [],
            "sheets": list(wb.sheetnames)}
    try:
        if "Summary" in wb.sheetnames:
            for row in wb["Summary"].iter_rows(values_only=True):
                if row and row[0] is not None:
                    info["summary"].append((str(row[0]), str(row[1]) if len(row) > 1 else ""))
        if "Status_bad" in wb.sheetnames:
            rows = list(wb["Status_bad"].iter_rows(max_row=60, values_only=True))
            info["status_bad"] = [
                [("" if c is None else str(c)) for c in r] for r in rows if r and any(r)
            ]
    finally:
        wb.close()
    return info


def read_change_log(log_path, tag=None):
    """Parse the companion .log for its changed-sensor lines - far cheaper than the xlsx."""
    entries = []
    if not log_path.is_file():
        return entries
    line_re = change_line_re(tag)
    for line in log_path.read_text(encoding="utf-8", errors="replace").splitlines():
        m = line_re.match(line.strip())
        if not m:
            continue
        d = m.groupdict()
        sev, why = classify(d["name"], d["delta"], d["max"])
        d.update(severity=sev, why=why)
        entries.append(d)
    entries.sort(key=lambda e: (SEVERITY_ORDER[e["severity"]], -int(e["count"]), e["name"]))
    return entries


def _esc(value):
    return html.escape(str(value))


def _render_source(info, entries, xlsx_name, raw_df=None, dq_findings=None, window=None):
    counts = {}
    for e in entries:
        counts[e["severity"]] = counts.get(e["severity"], 0) + 1
    for f in (dq_findings or []):
        counts[f["severity"]] = counts.get(f["severity"], 0) + 1

    chips = "".join(
        f'<span class="chip" style="background:{SEVERITY_COLOR[s]}">{s}: {counts[s]}</span>'
        for s in ("CRITICAL", "WARNING", "NOTICE", "INFO") if counts.get(s)
    ) or '<span class="chip" style="background:#666">no significant changes</span>'

    summary_rows = "".join(
        f"<tr><td>{_esc(k)}</td><td>{_esc(v)}</td></tr>"
        for k, v in (info["summary"] if info else [])
    ) or "<tr><td colspan=2>Summary sheet not available</td></tr>"

    status_rows = "".join(
        "<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in row) + "</tr>"
        for row in (info["status_bad"] if info else [])
    ) or "<tr><td>Status_bad sheet not available</td></tr>"

    change_rows = "".join(
        f'<tr class="sev-{e["severity"].lower()}">'
        f'<td><span class="sev" style="background:{SEVERITY_COLOR[e["severity"]]}">'
        f'{e["severity"]}</span></td>'
        f'<td class="name" title="{_esc(e["name"])}">{_esc(e["name"])}</td>'
        f'<td>{_esc(e["count"])}</td><td>{_esc(e["first"])}</td><td>{_esc(e["last"])}</td>'
        f'<td><b>{_esc(e["delta"])}</b></td><td>{_esc(e["min"])}</td><td>{_esc(e["max"])}</td>'
        f'<td class="why">{_esc(e["why"])}</td></tr>'
        for e in entries
    ) or '<tr><td colspan="9">No significant changes recorded.</td></tr>'

    sheet_count = len(info["sheets"]) if info else 0

    if window and window["dropped"]:
        window_html = (
            f'<div class="window">Analysis window: <b>{window["start"]:%Y-%m-%d %H:%M:%S}</b>'
            f' &rarr; <b>{window["end"]:%Y-%m-%d %H:%M:%S}</b> UTC &middot; '
            f'{window["sessions"]} sampling sessions found, only the last one is '
            f'analysed &middot; {window["dropped"]:,} of {window["total"]:,} rows from '
            f'earlier power states were excluded (sampling stops while the host is '
            f'down, so a gap &gt; {window["threshold_s"] / 60:.0f} min against the '
            f'{window["cadence_s"]:.1f} s cadence marks a power cycle). '
            f'All times are the CSV\'s own UTC stamps - no local-time conversion.</div>'
        )
    elif window:
        window_html = (
            f'<div class="window">Analysis window: <b>{window["start"]:%Y-%m-%d %H:%M:%S}</b>'
            f' &rarr; <b>{window["end"]:%Y-%m-%d %H:%M:%S}</b> UTC &middot; '
            f'one continuous sampling session, nothing excluded.</div>'
        )
    else:
        window_html = ""

    # Charts for the CRITICAL changed-sensor findings: the whole point is to see the
    # shape of the run, not just first/last numbers.
    critical = [e for e in entries if e["severity"] == "CRITICAL"]
    chart_blocks = []
    for e in critical:
        pairs = series_for(raw_df, e["name"]) if raw_df is not None else []
        chart_blocks.append(
            f'<div class="card">'
            f'<div class="cname">{_esc(e["name"])}</div>'
            f'<div class="cmeta">first <b>{_esc(e["first"])}</b> &rarr; last '
            f'<b>{_esc(e["last"])}</b> &middot; delta <b>{_esc(e["delta"])}</b> '
            f'&middot; min {_esc(e["min"])} &middot; max {_esc(e["max"])} '
            f'&middot; {_esc(e["count"])} spike(s)</div>'
            f'{svg_chart(pairs)}</div>'
        )

    # Textual CRITICAL findings (LinkStatus etc.) have no numeric series, so a
    # line chart is impossible - render the state history as a colour band.
    for f in (dq_findings or []):
        if f["severity"] != "CRITICAL":
            continue
        pairs = state_series_for(raw_df, f["sensor"]) if raw_df is not None else []
        chart_blocks.append(
            f'<div class="card">'
            f'<div class="cname">{_esc(f["sensor"])}</div>'
            f'<div class="cmeta">{_esc(f["kind"])} &middot; ends on '
            f'<b>{_esc(f["last"])}</b> &middot; {_esc(f["why"])}</div>'
            f'{svg_state_timeline(pairs)}</div>'
        )

    n_charts = len(chart_blocks)
    charts_html = (
        f'<h3>CRITICAL sensors, start &rarr; end ({n_charts})</h3>'
        + "".join(chart_blocks)
    ) if chart_blocks else ""

    dq_rows = "".join(
        f'<tr class="sev-{f["severity"].lower()}">'
        f'<td><span class="sev" style="background:{SEVERITY_COLOR[f["severity"]]}">'
        f'{f["severity"]}</span></td>'
        f'<td>{_esc(f["kind"])}</td>'
        f'<td class="name" title="{_esc(f["sensor"])}">{_esc(f["sensor"])}</td>'
        f'<td>{_esc(f["detail"])}</td><td>{_esc(f["last"])}</td>'
        f'<td>{_esc(f["count"]) if f["count"] else "-"}</td>'
        f'<td class="why">{_esc(f["why"])}</td></tr>'
        for f in (dq_findings or [])
    ) or '<tr><td colspan="7">No stale, null, 0xFF, bad-state or missing sensors found.</td></tr>'

    return f"""
<section>
  <h2>{_esc(xlsx_name)}</h2>
  <div class="meta">{chips}
    <a class="open" href="{_esc(xlsx_name)}">Open workbook ({sheet_count} sheets) &#8599;</a>
  </div>
  <h3>Summary</h3>
  <table class="kv">{summary_rows}</table>
  <h3>Status anomalies (Status_bad)</h3>
  <table class="kv">{status_rows}</table>
  {charts_html}
  <h3>Data quality &amp; sensor state ({len(dq_findings or [])})</h3>
  {window_html}
  <div class="note">Checks the raw CSV for things the workbook cannot show:
  every non-numeric reading is coerced to NaN there, and
  <code>MetricValueStale</code> is typically not read by the workbook generator.</div>
  <table class="grid">
    <tr><th>Severity</th><th>Check</th><th>Sensor</th><th>Observed</th>
        <th>Last</th><th>Samples</th><th>Why it is listed</th></tr>
    {dq_rows}
  </table>
  <h3>Significant changes ({len(entries)})</h3>
  <table class="grid">
    <tr><th>Severity</th><th>Sensor</th><th>#</th><th>First</th><th>Last</th>
        <th>Delta</th><th>Min</th><th>Max</th><th>Why it is listed</th></tr>
    {change_rows}
  </table>
</section>
"""


def build_digest(scan_folder, output=None, xlsx_glob=None, workbook_suffix=None,
                 change_tag=None):
    """Write the digest HTML for every workbook matching ``xlsx_glob``.

    Returns the output path, or None when no workbook was found (a normal case,
    e.g. a run without sensor collection).
    """
    scan_folder = Path(scan_folder)
    xlsx_glob = xlsx_glob or SENSOR_XLSX_GLOB
    books = sorted(p for p in scan_folder.glob(xlsx_glob)
                   if p.is_file() and not p.name.startswith("~$"))
    if not books:
        logger.info("No %s in %s - skipping sensor digest.", xlsx_glob, scan_folder)
        return None

    sections, total_critical = [], 0
    for xlsx in books:
        info = read_summary(xlsx)
        entries = read_change_log(xlsx.with_suffix(".log"), change_tag)
        raw_csv = raw_csv_for(xlsx, workbook_suffix)
        raw_df = load_raw_csv(raw_csv) if raw_csv else None
        window = detect_active_window(raw_df)
        if window and window["dropped"]:
            raw_df = raw_df[raw_df["Timestamp"] >= window["start"]]
            logger.info("%s: analysing the last of %d sessions (%s -> %s), "
                        "dropped %d rows from earlier power states",
                        xlsx.name, window["sessions"], window["start"],
                        window["end"], window["dropped"])
        if raw_csv is None:
            logger.info("No raw CSV next to %s - charts/data-quality skipped.", xlsx.name)
        dq = check_data_quality(raw_df) if raw_df is not None else []
        entries = demote_changes_by_peers(entries, raw_df)
        total_critical += sum(1 for e in entries if e["severity"] == "CRITICAL")
        total_critical += sum(1 for f in dq if f["severity"] == "CRITICAL")
        sections.append(_render_source(info, entries, xlsx.name, raw_df, dq, window))

    banner = brand.brand_banner("Sensor Digest", str(scan_folder)) if brand else ""
    footer = brand.brand_footer("Sensor Digest") if brand else ""
    favicon = brand.favicon_tag() if brand else ""
    brand_css = brand.brand_css() if brand else ""

    out = Path(output) if output else scan_folder / OUTPUT_NAME
    out.write_text(f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Sensor XLSX Digest</title>
{favicon}
<style>{brand_css}
body{{font-family:Segoe UI,Arial,sans-serif;margin:0;background:#f7f7f9;color:#222}}
.wrap{{max-width:1500px;margin:20px auto;padding:0 16px}}
section{{background:#fff;border:1px solid #ddd;border-radius:6px;
padding:14px 18px;margin-bottom:18px}}
h2{{font-size:18px;margin:0 0 10px}} h3{{font-size:14px;margin:16px 0 6px;color:#444}}
.meta{{margin-bottom:8px}}
.chip{{color:#fff;font-size:12px;padding:2px 9px;border-radius:10px;margin-right:6px}}
.open{{font-size:13px;color:#1a5fb4;text-decoration:none;font-weight:600}}
table{{border-collapse:collapse;width:100%;font-size:12px}}
.kv td{{border:1px solid #e2e2e2;padding:4px 8px}}
.kv td:first-child{{width:220px;color:#555;background:#fafafa}}
.grid th{{background:#2c3446;color:#fff;padding:5px 7px;text-align:left;
position:sticky;top:0}}
.grid td{{border-bottom:1px solid #eee;padding:4px 7px;vertical-align:top}}
.grid tr.sev-critical{{background:#fff2f2}} .grid tr.sev-warning{{background:#fff8ef}}
.sev{{color:#fff;font-size:11px;padding:1px 6px;border-radius:8px}}
.name{{font-family:Consolas,monospace;max-width:620px;word-break:break-all}}
.why{{color:#666}}
.note{{color:#666;font-size:12px;margin:0 0 8px}}
.card{{border:1px solid #f0c9c9;border-left:4px solid #c62828;border-radius:5px;
padding:8px 12px;margin:0 0 12px;background:#fffdfd}}
.cname{{font-family:Consolas,monospace;font-size:12px;word-break:break-all;color:#333}}
.cmeta{{font-size:12px;color:#666;margin:3px 0 5px}}
.spark{{display:block}}
.nochart{{color:#999;font-size:12px;padding:8px 0}}
.window{{background:#eef4fb;border-left:4px solid #1565c0;border-radius:4px;
padding:8px 12px;margin:0 0 8px;font-size:12px;color:#334}}
</style></head><body>
<div class="wrap">
{banner}
<div style="color:#666;font-size:13px;margin-bottom:14px">
{_esc(scan_folder)} &middot; generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
&middot; <b style="color:{SEVERITY_COLOR['CRITICAL']}">{total_critical} CRITICAL</b>
</div>
{''.join(sections)}
{footer}
</div></body></html>
""", encoding="utf-8")
    logger.info("Wrote sensor digest (%d workbooks, %d critical): %s",
                len(books), total_critical, out)
    return out


def main():
    parser = argparse.ArgumentParser(
        description="Digest sensor-analysis workbooks (*.xlsx) into one HTML page")
    parser.add_argument("--scan_folder", required=True,
                        help="Folder containing the workbooks (and optional .log / .csv companions).")
    parser.add_argument("--output", help=f"Output HTML (default: <scan_folder>/{OUTPUT_NAME}).")
    parser.add_argument("--xlsx_glob", default=SENSOR_XLSX_GLOB,
                        help=f"Workbook glob relative to --scan_folder (default: {SENSOR_XLSX_GLOB}; "
                             "env: LOG_ANALYZER_SENSOR_XLSX_GLOB).")
    parser.add_argument("--workbook_suffix", default=WORKBOOK_SUFFIX,
                        help="Workbook-name suffix replaced by .csv to locate the raw collector CSV "
                             f"(default: {WORKBOOK_SUFFIX}; env: LOG_ANALYZER_SENSOR_WORKBOOK_SUFFIX).")
    parser.add_argument("--change_tag", default=CHANGE_LINE_TAG,
                        help="Tag of changed-sensor lines in the companion .log "
                             f"(default: {CHANGE_LINE_TAG}; env: LOG_ANALYZER_SENSOR_CHANGE_TAG).")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    # Having no sensor workbook is a normal case (e.g. a collect-only run), so
    # it must not look like a failure to the caller.
    build_digest(args.scan_folder, args.output, xlsx_glob=args.xlsx_glob,
                 workbook_suffix=args.workbook_suffix, change_tag=args.change_tag)
    return 0


if __name__ == "__main__":
    sys.exit(main())
